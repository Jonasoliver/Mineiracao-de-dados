# Gera o documento Word com as solucoes dos 10 exercicios usando python-docx
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = "/sessions/tender-fervent-curie/mnt/outputs"
ARIAL = "Arial"
MONO = "Consolas"

def set_cell_shading(paragraph, fill_hex):
    """Aplica sombreamento de fundo a um paragrafo (parametro pPr/shd)."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)

def add_para(doc, text, *, bold=False, italic=False, size=11, align=None, font=ARIAL, after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p

def add_heading(doc, text, level, color="000000"):
    sizes = {1: 18, 2: 14, 3: 12}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = ARIAL
    run.font.size = Pt(sizes.get(level, 12))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(color)
    # Marca como heading para indicar nivel
    p.style = doc.styles[f'Heading {level}']
    # Substitui a fonte do estilo
    run.font.name = ARIAL
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(sizes.get(level, 12))
    return p

def add_code_block(doc, code, fill="F2F2F2"):
    for line in code.split('\n'):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        run = p.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(9)
        # Garante que a fonte apareca em editores asiaticos / fallback
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), MONO)
        rFonts.set(qn('w:hAnsi'), MONO)
        rFonts.set(qn('w:cs'), MONO)
        rPr.append(rFonts)
        set_cell_shading(p, fill)

def le(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()

# Conteudo de cada exercicio: (titulo, enunciado, ideia, arquivo, saida)
EXERCICIOS = [
    (
        "Exercicio 1 - Extracao de dominio de e-mail e flag provedor x empresarial",
        "Crie um script em Python que processe uma lista de e-mails, extraia o dominio de cada endereco e gere uma variavel binaria (Flag) que identifique se o registro pertence a um provedor comum ou a uma infraestrutura empresarial (ex.: terminando em \".com.br\").",
        "A logica usa string.split('@')[-1] para isolar o dominio e compara com um conjunto de provedores publicos conhecidos (Gmail, Hotmail, Yahoo etc.). Tudo que nao bate com essa lista e tratado como dominio empresarial. Geramos duas flags binarias complementares para alimentar o modelo: flag_provedor_comum e flag_empresarial.",
        f"{OUT_DIR}/exercicio_01_email_dominio.py",
        """                  email              dominio  flag_provedor_comum  flag_empresarial
   joao.silva@gmail.com            gmail.com                    1                 0
   maria@empresa.com.br       empresa.com.br                    0                 1
     contato@startup.io           startup.io                    0                 1
       user@hotmail.com          hotmail.com                    1                 0
rh@bancodobrasil.com.br bancodobrasil.com.br                    0                 1
       vendas@yahoo.com            yahoo.com                    1                 0""",
    ),
    (
        "Exercicio 2 - Decomposicao de timestamp e flag de fim de semana",
        "A partir de uma lista de tempo (timestamps), desenvolva um codigo que decomponha a data para extrair o mes e o dia da semana, gerando uma Flag adicional para identificar automaticamente acessos realizados no Final de Semana.",
        "Convertemos a coluna para datetime e usamos os acessores .dt.month e .dt.dayofweek do pandas. O dayofweek devolve 0 (segunda) ate 6 (domingo); a flag de fim de semana e simplesmente isin([5, 6]). O mapeamento numerico do dia para o nome em portugues facilita a interpretabilidade do modelo.",
        f"{OUT_DIR}/exercicio_02_timestamp_fim_de_semana.py",
        """          timestamp  mes  dia_semana_num dia_semana  flag_fim_de_semana
2026-05-04 09:15:00    5               0    Segunda                   0
2026-05-09 22:30:00    5               5     Sabado                   1
2026-05-10 11:00:00    5               6    Domingo                   1
2026-12-25 08:00:00   12               4      Sexta                   0
2026-07-15 14:20:00    7               2     Quarta                   0""",
    ),
    (
        "Exercicio 3 - Flag de feriado / evento comercial",
        "Desenvolva um programa que receba uma data de transacao e verifique se ela coincide com um feriado ou evento comercial (ex.: Natal ou Black Friday), gerando uma variavel binaria para que o modelo de IA reconheca o pico de vendas como um evento esperado e nao uma anomalia.",
        "Mantemos um dicionario de feriados fixos (mes, dia) e uma funcao para datas variaveis como Black Friday (ultima sexta de novembro). A flag_evento sinaliza para o modelo que aquele pico esta calendarizado, evitando que ele seja sinalizado como outlier por algoritmos de deteccao de anomalia.",
        f"{OUT_DIR}/exercicio_03_feriados_eventos.py",
        """      data        evento  flag_evento
2026-12-25         Natal            1
2026-11-27  Black Friday            1
2026-07-04                          0
2026-05-10                          0
2026-09-07 Independencia            1""",
    ),
    (
        "Exercicio 4 - Resumo RFM (Recencia, Frequencia, Monetario)",
        "Com base em um historico de transacoes contendo ID do cliente, data e valor, utilize tecnicas de agrupamento para transformar o log transacional em um resumo de perfil unico que contenha a Recencia (dias desde a ultima compra), a Frequencia (contagem de pedidos) e o Valor Monetario (soma total), conforme o modelo RFM.",
        "O groupby por id_cliente reduz o log (1 linha por compra) a 1 linha por cliente. Recencia e a diferenca em dias entre a data de referencia e a maior data de compra; Frequencia e a contagem de transacoes; Monetario e a soma do valor. Esse formato e a base para segmentacoes (Champions, At Risk etc.) e para algoritmos como K-Means.",
        f"{OUT_DIR}/exercicio_04_rfm.py",
        """ id_cliente  recencia  frequencia  monetario
          1         5           3      470.5
          2        84           2      730.0
          3         0           3      190.0""",
    ),
    (
        "Exercicio 5 - Delta e Evolucao Percentual de sensor",
        "Desenvolva um script que monitore a leitura de um sensor em dois momentos distintos, calculando o Delta (variacao bruta) e a Evolucao Percentual entre eles, sinalizando se a tendencia atual e de crescimento ou queda para contextualizar o dado para o algoritmo.",
        "Delta = leitura_t2 - leitura_t1 da o tamanho absoluto da variacao. A evolucao percentual normaliza esse delta pela base inicial (delta / leitura_t1 x 100), tornando comparaveis sensores em escalas diferentes. Tratamos o caso leitura_t1 == 0 para evitar divisao por zero. A coluna 'tendencia' (Crescimento, Queda, Estavel) e uma feature categorica derivada que da contexto direto ao algoritmo.",
        f"{OUT_DIR}/exercicio_05_delta_sensor.py",
        """      t1       t2      delta   evolucao_%      tendencia
   100.0    115.0       15.0         15.0    Crescimento
    80.5     80.5        0.0          0.0        Estavel
   200.0    150.0      -50.0        -25.0          Queda
    50.0     75.0       25.0         50.0    Crescimento""",
    ),
    (
        "Exercicio 6 - Razao de Comprometimento (Divida/Renda)",
        "Em um cenario de analise de credito, crie uma nova feature baseada na interacao entre atributos que calcule a Razao de Comprometimento (Divida total dividida pela Renda mensal), justificando por que essa proporcao e mais valiosa para a IA do que as variaveis isoladas.",
        "A razao e uma feature de interacao: combina dois atributos num indicador que carrega significado financeiro (% da renda comprometida). Diferente de divida e renda isoladas, ela e invariante a escala absoluta - 4.000 de divida nao significa nada sem o salario de referencia. Tambem permite faixas de risco diretas (Baixo, Moderado, Alto, Critico) que o modelo aprende com menos dados. Em resumo: a razao consolida o sinal em uma unica feature mais informativa, mais resistente a outliers absolutos e com semantica de negocio direta para o algoritmo.",
        f"{OUT_DIR}/exercicio_06_razao_comprometimento.py",
        """ id  divida  renda  razao_comp             faixa_risco
  1    1500   5000        0.30                Moderado
  2    4000   5000        0.80                 Critico
  3     800   5000        0.16                   Baixo
  4    1500   2000        0.75                    Alto
  5    1500      0         inf Indefinido (renda zero)""",
    ),
    (
        "Exercicio 7 - Indicador de Omissao (flag de NaN antes do preenchimento)",
        "Desenvolva um tratamento para campos nulos (NaN) que, antes de realizar qualquer limpeza ou preenchimento, crie um Indicador de Omissao (Flag de erro) para preservar a informacao de falha de leitura ou escolha do usuario, preenchendo a lacuna original apenas em seguida.",
        "Antes de imputar (mediana / media / zero), criamos uma coluna binaria <coluna>_foi_omisso = 1 onde havia NaN. Isso conserva o sinal de 'o usuario nao informou' ou 'sensor falhou', informacao que muitas vezes e preditiva por si so. Sem essa flag, depois do fillna o modelo passa a tratar valor imputado e valor real como indistinguiveis.",
        f"{OUT_DIR}/exercicio_07_indicador_omissao.py",
        """ANTES:
 id  renda  idade
  1 3500.0   25.0
  2    NaN   32.0
  3 8000.0    NaN
  4 5200.0   41.0
  5    NaN   29.0

DEPOIS (com flag de omissao):
 id  renda  idade  renda_foi_omisso  idade_foi_omisso
  1 3500.0   25.0                 0                 0
  2 5200.0   32.0                 1                 0
  3 8000.0   30.5                 0                 1
  4 5200.0   41.0                 0                 0
  5 5200.0   29.0                 1                 0""",
    ),
    (
        "Exercicio 8 - Ranking Interno por volume de vendas",
        "Construa um programa que receba o volume de vendas de diversos produtos e gere uma nova coluna de Ranking Interno, posicionando cada item em relacao aos seus pares para que o modelo entenda a forca relativa de cada registro dentro do grupo.",
        "O metodo .rank() do pandas retorna a posicao de cada valor; com method='dense' nao ha 'pulos' de posicao em empates. A versao com groupby gera ranking por categoria, separando os pares relevantes (Eletronicos x Eletronicos, Vestuario x Vestuario). Tambem geramos um percentil (.rank(pct=True)) que normaliza a posicao para [0,1] - util para modelos que precisam de escala uniforme.",
        f"{OUT_DIR}/exercicio_08_ranking_interno.py",
        """Ranking GLOBAL:
produto   categoria  vendas  ranking_interno  percentil
      C Eletronicos    1500                1      1.000
      A Eletronicos    1200                2      0.833
      B Eletronicos     900                3      0.667
      E   Vestuario     780                4      0.500
      D   Vestuario     450                5      0.333
      F   Vestuario     300                6      0.167

Ranking POR CATEGORIA:
produto   categoria  vendas  ranking_interno  percentil
      C Eletronicos    1500                1      1.000
      A Eletronicos    1200                2      0.667
      B Eletronicos     900                3      0.333
      E   Vestuario     780                1      1.000
      D   Vestuario     450                2      0.667
      F   Vestuario     300                3      0.333""",
    ),
    (
        "Exercicio 9 - Parser de horario para Turnos",
        "Desenvolva um parser que fatie a variavel de horario de um log (00h as 23h) em categorias de Turnos (ex.: Madrugada, Comercial, Noite), transformando um dado numerico continuo em blocos de comportamento humano mais simples de processar.",
        "Convertemos o timestamp em hora cheia e classificamos em quatro turnos: Madrugada (0-5h), Manha (6-11h), Comercial (12-17h) e Noite (18-23h). Esse binning reduz a cardinalidade da variavel de 24 valores para 4 categorias com semantica de comportamento humano - acesso de madrugada tem perfil de risco diferente de acesso comercial, e isso fica explicito na feature.",
        f"{OUT_DIR}/exercicio_09_turnos.py",
        """          timestamp  hora     turno
2026-05-10 02:15:00     2 Madrugada
2026-05-10 08:45:00     8     Manha
2026-05-10 13:30:00    13 Comercial
2026-05-10 19:50:00    19     Noite
2026-05-10 23:10:00    23     Noite
2026-05-10 06:00:00     6     Manha
2026-05-10 17:59:00    17 Comercial""",
    ),
    (
        "Exercicio 10 - Delta do Delta (aceleracao) e alerta de tendencia exponencial",
        "Crie um script que calcule a diferenca entre as ultimas tres leituras de um sensor para identificar a aceleracao da variavel (o Delta do Delta), disparando um alerta binario sempre que o ritmo de crescimento indicar uma tendencia de aumento exponencial.",
        "Com tres pontos (l1, l2, l3) calculamos delta_1 = l2 - l1 e delta_2 = l3 - l2. A aceleracao e a diferenca entre os dois deltas (delta_2 - delta_1) - matematicamente a segunda diferenca discreta, equivalente a derivada segunda. O alerta dispara apenas quando ambos os deltas e a aceleracao sao positivos, padrao tipico de crescimento exponencial. Linear (5,5), desacelerando (8,4) e queda nao disparam.",
        f"{OUT_DIR}/exercicio_10_delta_do_delta.py",
        """Crescimento exponencial        -> {'leituras': [10, 12, 16], 'delta_1': 2, 'delta_2': 4, 'aceleracao': 2, 'alerta_exponencial': 1}
Crescimento linear             -> {'leituras': [10, 15, 20], 'delta_1': 5, 'delta_2': 5, 'aceleracao': 0, 'alerta_exponencial': 0}
Crescimento desacelerando      -> {'leituras': [10, 18, 22], 'delta_1': 8, 'delta_2': 4, 'aceleracao': -4, 'alerta_exponencial': 0}
Queda                          -> {'leituras': [50, 40, 30], 'delta_1': -10, 'delta_2': -10, 'aceleracao': 0, 'alerta_exponencial': 0}
Estavel                        -> {'leituras': [10, 10, 10], 'delta_1': 0, 'delta_2': 0, 'aceleracao': 0, 'alerta_exponencial': 0}""",
    ),
]

doc = Document()

# Margens
for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

# Estilo padrao
normal = doc.styles['Normal']
normal.font.name = ARIAL
normal.font.size = Pt(11)

# Capa
add_para(doc, "Mineracao de Dados", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
add_para(doc, "Exercicios - Unidade 10", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
add_para(doc, "Resolucao em Python", italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

add_para(
    doc,
    "Este documento apresenta a resolucao dos 10 exercicios da Unidade 10 de Mineracao de Dados. "
    "Para cada exercicio sao apresentados: o enunciado original, a ideia da solucao, o codigo Python "
    "completo e a saida obtida ao executar o script.",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    after=12,
)

# Exercicios
for titulo, enunciado, ideia, arquivo, saida in EXERCICIOS:
    add_heading(doc, titulo, level=2, color="1F3864")

    add_heading(doc, "Enunciado", level=3, color="2E75B6")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(enunciado)
    run.font.name = ARIAL
    run.font.size = Pt(11)
    run.italic = True

    add_heading(doc, "Ideia da solucao", level=3, color="2E75B6")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(ideia)
    run.font.name = ARIAL
    run.font.size = Pt(11)

    add_heading(doc, "Codigo Python", level=3, color="2E75B6")
    add_code_block(doc, le(arquivo), fill="F2F2F2")

    # Espaco
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)

    add_heading(doc, "Saida obtida", level=3, color="2E75B6")
    add_code_block(doc, saida, fill="FFF8E1")

    # Espaco maior antes do proximo exercicio
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)

out_path = f"{OUT_DIR}/Exercicios_Unidade10_Resolucao.docx"
doc.save(out_path)
print(f"OK: {out_path}")
